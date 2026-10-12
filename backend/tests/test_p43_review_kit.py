"""M5.5 P4-3.2 人工审核工作流（``scripts/p43_review_kit.py``）测试。

分三层：
  1. **纯逻辑**：``collect_decisions`` 的校验规则（E00–E20），不碰文件；
  2. **xlsx 往返**：build → 人工填写 → read_checklist → collect；
  3. **CLI 端到端**：子进程跑 build / collect，验证退出码与落盘产物。

刻意不联网、不碰数据库。
"""

from __future__ import annotations

import ast
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
KIT_PATH = ROOT / "scripts" / "p43_review_kit.py"
GATES_PATH = ROOT / "scripts" / "p43_gates.py"

FIXED_TS = "2026-10-12T00:00:00+00:00"
CONCEPT = {"id": 11, "name": "顺序表", "code": "DS.LINEAR.SEQUENTIAL"}
GOOD_ANALYSIS = (
    "顺序表插入的最坏情况是把新元素插到表头，此时原有的 n 个元素必须整体后移一位，"
    "因此需要移动 n 个元素；这也正是顺序表插入代价高于链表的根本原因。"
)


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


assert GATES_PATH.is_file(), f"缺少门禁模块：{GATES_PATH}"
kit = load_module("p43_review_kit", KIT_PATH)
gates = kit.gates


# ---------------------------------------------------------------------------
# fixture 构造
# ---------------------------------------------------------------------------
def make_question(**overrides):
    question = {
        "subject": "数据结构",
        "knowledge_point": CONCEPT["name"],
        "question_type": "choice",
        "stem": "对一个长度为 n 的顺序表执行插入操作，最坏情况下需要移动多少个元素？",
        "options": ["A. n 个", "B. n-1 个", "C. n/2 个", "D. 1 个"],
        "answer": "A",
        "analysis": GOOD_ANALYSIS,
        "difficulty": "easy",
    }
    question.update(overrides)
    return question


def make_entry(questions, concept=None):
    concept = concept or CONCEPT
    return {
        "concept": {**concept, "order": 1},
        "question_type": "choice",
        "error": None,
        "raw_analysis": {index: item["analysis"] for index, item in enumerate(questions)},
        "questions": questions,
    }


def default_manifest():
    good_one = make_question()
    good_two = make_question(
        stem="在顺序表中删除第 i 个元素时，最坏情况下需要移动多少个元素？",
        answer="B",
        difficulty="medium",
    )
    failing = make_question(
        stem="关于顺序表的随机访问特性，下列说法正确的是？",
        options=["A. 可以按序号直接访问", "B. 只能顺序访问", "C. 无法访问数据", "D. 需要哈希"],
        answer="D",
        analysis="顺序表可以随机访问。",  # <60 字 → G3 判错
    )
    return gates.assemble_manifest(
        subject="数据结构",
        entries=[make_entry([good_one, good_two]), make_entry([failing])],
        generated_at=FIXED_TS,
        batch=1,
    )


def question_id_at(manifest, index):
    return manifest["questions"][index]["question_id"]


def write_manifest(tmp_path: Path, manifest) -> Path:
    path = tmp_path / "manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def build_checklist(tmp_path: Path, manifest, manifest_path: Path, *, lock: bool = True) -> Path:
    path = tmp_path / "checklist.xlsx"
    workbook = kit.build_workbook(
        manifest,
        lock=lock,
        built_at=FIXED_TS,
        manifest_sha=kit.sha256_file(manifest_path),
        manifest_name=manifest_path.name,
    )
    workbook.save(path)
    return path


@pytest.fixture
def batch(tmp_path):
    manifest = default_manifest()
    manifest_path = write_manifest(tmp_path, manifest)
    checklist_path = build_checklist(tmp_path, manifest, manifest_path)
    return {
        "manifest": manifest,
        "manifest_path": manifest_path,
        "checklist_path": checklist_path,
        "tmp_path": tmp_path,
    }


def header_index(path: Path) -> dict[str, int]:
    workbook = load_workbook(path)
    headers = [cell.value for cell in workbook[kit.CHECKLIST_SHEET][1]]
    return {header: index + 1 for index, header in enumerate(headers) if header}


def edit_rows(path: Path, per_question: dict[str, dict]):
    """按 question_id 就地改单元格（模拟人工在 Excel 里编辑）。

    注意：openpyxl 的 ``cell(value=None)`` 是**空操作**，清空单元格必须显式赋 ``.value``。
    """
    workbook = load_workbook(path)
    sheet = workbook[kit.CHECKLIST_SHEET]
    index = header_index(path)
    for row in range(2, sheet.max_row + 1):
        qid = sheet.cell(row=row, column=index[kit.COL_QID]).value
        for header, value in per_question.get(qid, {}).items():
            sheet.cell(row=row, column=index[header]).value = value
    workbook.save(path)


def append_row(path: Path, values: dict):
    workbook = load_workbook(path)
    sheet = workbook[kit.CHECKLIST_SHEET]
    index = header_index(path)
    row = sheet.max_row + 1
    for header, value in values.items():
        sheet.cell(row=row, column=index[header]).value = value
    workbook.save(path)


def run_collect(checklist_path: Path, manifest_path: Path, *, existing=None, collected_at=FIXED_TS):
    headers, rows, meta = kit.read_checklist(checklist_path)
    manifest = kit.load_manifest(manifest_path)
    return kit.collect_decisions(
        manifest,
        headers,
        rows,
        expected_sha=meta.get("manifest SHA-256"),
        actual_sha=kit.sha256_file(manifest_path),
        existing=existing or [],
        collected_at=collected_at,
    )


def codes(result) -> list[str]:
    return [item["code"] for item in result["errors"]]


# ---------------------------------------------------------------------------
# 构表
# ---------------------------------------------------------------------------
def test_build_creates_three_sheets_and_expected_headers(batch):
    workbook = load_workbook(batch["checklist_path"])
    assert workbook.sheetnames == [kit.CHECKLIST_SHEET, kit.GUIDE_SHEET, kit.BATCH_SHEET]
    headers = [cell.value for cell in workbook[kit.CHECKLIST_SHEET][1]]
    assert headers[:4] == [kit.COL_QID, kit.COL_CONCEPT, kit.COL_CODE, kit.COL_GATE]
    assert headers.count("选项A") == 1 and headers.count("选项D") == 1
    assert headers[-2:] == [kit.COL_DECISION, kit.COL_NOTE]


def test_build_option_columns_cover_batch_max_and_floor_at_four(tmp_path):
    manifest = default_manifest()
    manifest_path = write_manifest(tmp_path, manifest)
    path = build_checklist(tmp_path, manifest, manifest_path)
    headers = [cell.value for cell in load_workbook(path)[kit.CHECKLIST_SHEET][1]]
    option_headers = [item for item in headers if item.startswith("选项")]
    assert option_headers == ["选项A", "选项B", "选项C", "选项D"]

    two_option = make_question(options=["A. 是", "B. 否"], answer="A")
    small = gates.assemble_manifest(
        subject="数据结构",
        entries=[make_entry([two_option])],
        generated_at=FIXED_TS,
        batch=2,
    )
    small_path = build_checklist(tmp_path, small, write_manifest(tmp_path, small))
    headers = [cell.value for cell in load_workbook(small_path)[kit.CHECKLIST_SHEET][1]]
    assert [item for item in headers if item.startswith("选项")] == [
        "选项A",
        "选项B",
        "选项C",
        "选项D",
    ]


def test_build_writes_gate_verdict_and_reasons(batch):
    workbook = load_workbook(batch["checklist_path"])
    sheet = workbook[kit.CHECKLIST_SHEET]
    index = header_index(batch["checklist_path"])
    gates_seen = [
        sheet.cell(row=row, column=index[kit.COL_GATE]).value
        for row in range(2, sheet.max_row + 1)
    ]
    assert gates_seen == ["PASS", "PASS", "FAIL"]
    fail_row = sheet.max_row
    reason = sheet.cell(row=fail_row, column=index[kit.COL_GATE_REASON]).value
    assert "G3" in reason


def test_build_lock_defaults_on_and_can_be_disabled(tmp_path):
    manifest = default_manifest()
    manifest_path = write_manifest(tmp_path, manifest)
    locked = load_workbook(build_checklist(tmp_path, manifest, manifest_path))
    assert locked[kit.CHECKLIST_SHEET].protection.sheet is True

    unlocked = load_workbook(build_checklist(tmp_path, manifest, manifest_path, lock=False))
    assert unlocked[kit.CHECKLIST_SHEET].protection.sheet is not True


def test_build_locks_metadata_and_unlocks_content(batch):
    workbook = load_workbook(batch["checklist_path"])
    sheet = workbook[kit.CHECKLIST_SHEET]
    index = header_index(batch["checklist_path"])
    assert sheet.cell(row=2, column=index[kit.COL_QID]).protection.locked is True
    assert sheet.cell(row=2, column=index[kit.COL_GATE]).protection.locked is True
    assert sheet.cell(row=2, column=index[kit.COL_STEM]).protection.locked is False
    assert sheet.cell(row=2, column=index["选项A"]).protection.locked is False
    assert sheet.cell(row=2, column=index[kit.COL_DECISION]).protection.locked is False


def test_build_adds_decision_dropdown(batch):
    workbook = load_workbook(batch["checklist_path"])
    sheet = workbook[kit.CHECKLIST_SHEET]
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    validation = validations[0]
    assert validation.type == "list"
    assert all(label in validation.formula1 for label in kit.DECISION_LABELS)
    assert "$" not in validation.formula1


def test_build_rejects_empty_manifest(tmp_path):
    path = tmp_path / "empty.json"
    path.write_text(json.dumps({"subject": "数据结构", "questions": []}), encoding="utf-8")
    with pytest.raises(kit.ReviewKitError):
        kit.load_manifest(path)


def test_build_is_byte_deterministic(tmp_path):
    manifest = default_manifest()
    manifest_path = write_manifest(tmp_path, manifest)
    first = build_checklist(tmp_path, manifest, manifest_path)
    second = tmp_path / "checklist2.xlsx"
    kit.build_workbook(
        manifest,
        built_at=FIXED_TS,
        manifest_sha=kit.sha256_file(manifest_path),
        manifest_name=manifest_path.name,
    ).save(second)
    assert first.read_bytes() == second.read_bytes()


# ---------------------------------------------------------------------------
# 回收：正常路径
# ---------------------------------------------------------------------------
def test_collect_clean_when_all_passing_approved(batch):
    ids = [question_id_at(batch["manifest"], index) for index in range(3)]
    edit_rows(
        batch["checklist_path"],
        {
            ids[0]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[1]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[2]: {kit.COL_DECISION: kit.DECISION_REJECT},
        },
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert result["errors"] == []
    assert result["summary"]["approve"] == 2
    assert result["summary"]["reject"] == 1
    assert result["summary"]["blocking"] == 0
    assert result["summary"]["ready_for_apply"] == 2
    approved = next(item for item in result["decisions"] if item["decision"] == "approve")
    assert approved["changed_fields"] == []
    assert approved["apply_hint"].startswith("submit_review")
    assert approved["final"]["stem"] == approved["original"]["stem"]


def test_collect_records_revision_diff(batch):
    qid = question_id_at(batch["manifest"], 0)
    new_stem = "顺序表在表头位置插入一个新元素时，最坏情况下需要移动多少个既有元素？"
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_STEM: new_stem}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    record = result["decisions"][0]
    assert record["decision"] == "revise"
    assert record["changed_fields"] == ["stem"]
    assert record["final"]["stem"] == new_stem
    assert record["revision_diff"][0]["before"] != record["revision_diff"][0]["after"]
    assert record["ready_for_apply"] is True
    assert record["errors"] == []


# ---------------------------------------------------------------------------
# 回收：校验规则
# ---------------------------------------------------------------------------
def test_collect_flags_unreviewed_rows(batch):
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E01" in codes(result)
    assert result["summary"]["unreviewed"] == 3
    assert result["summary"]["ready_for_apply"] == 0


def test_collect_flags_invalid_decision_label(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(batch["checklist_path"], {qid: {kit.COL_DECISION: "同意"}})
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E02" in codes(result)


def test_collect_requires_note_to_override_failed_gate(batch):
    failed = question_id_at(batch["manifest"], 2)
    edit_rows(batch["checklist_path"], {failed: {kit.COL_DECISION: kit.DECISION_APPROVE}})
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E03" in codes(result)
    record = next(item for item in result["decisions"] if item["question_id"] == failed)
    assert record["gate_passed"] is False
    assert record["ready_for_apply"] is False


def test_collect_allows_override_with_note_and_records_gates(batch):
    ids = [question_id_at(batch["manifest"], index) for index in range(3)]
    edit_rows(
        batch["checklist_path"],
        {
            ids[0]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[1]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[2]: {kit.COL_DECISION: kit.DECISION_APPROVE, kit.COL_NOTE: "G3 解析略短但结论正确，接受"},
        },
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert result["errors"] == []
    override = next(item for item in result["decisions"] if item["question_id"] == ids[2])
    assert override["overridden_gates"] == ["G3"]
    assert override["ready_for_apply"] is True


def test_collect_flags_revise_without_any_change(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(batch["checklist_path"], {qid: {kit.COL_DECISION: kit.DECISION_REVISE}})
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E05" in codes(result)


def test_collect_flags_approve_with_edits(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_APPROVE, kit.COL_ANALYSIS: GOOD_ANALYSIS + "补充说明。"}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E04" in codes(result)


def test_collect_flags_reject_with_edits(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REJECT, kit.COL_STEM: "改过的题干但判定为驳回"}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E09" in codes(result)


def test_collect_flags_broken_revision_answer_letter(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_ANSWER: "F"}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E06" in codes(result)
    record = result["decisions"][0]
    assert record["revised_gate_check"]["gates"]["G2"]["passed"] is False


def test_collect_flags_invalid_difficulty_after_revision(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_DIFFICULTY: "超难"}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E06" in codes(result)


def test_collect_warns_but_does_not_block_on_short_revised_analysis(batch):
    ids = [question_id_at(batch["manifest"], index) for index in range(3)]
    edit_rows(
        batch["checklist_path"],
        {
            ids[0]: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_ANALYSIS: "解析改短了。"},
            ids[1]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[2]: {kit.COL_DECISION: kit.DECISION_REJECT},
        },
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E07" in codes(result)
    assert result["summary"]["blocking"] == 0
    assert result["summary"]["warnings"] >= 1
    record = result["decisions"][0]
    assert "E07" not in [item["code"] for item in record["errors"]]
    assert "E07" in [item["code"] for item in record["warnings"]]
    assert record["ready_for_apply"] is True


def test_collect_flags_option_gap(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REJECT, "选项B": None}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E08" in codes(result)


# ---------------------------------------------------------------------------
# 回收：清单与 manifest 的一致性
# ---------------------------------------------------------------------------
def test_collect_flags_unknown_and_missing_ids(batch):
    known = question_id_at(batch["manifest"], 0)
    append_row(batch["checklist_path"], {kit.COL_QID: "P43-01-999"})
    edit_rows(batch["checklist_path"], {known: {kit.COL_QID: None}})
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    found = codes(result)
    assert "E10" in found
    assert "E13" in found
    assert "E12" in found


def test_collect_detects_manifest_sha_mismatch(batch):
    other = default_manifest()
    other["questions"][0]["stem"] = "换了内容的题干，用于制造指纹不一致。"
    other_path = write_manifest(batch["tmp_path"], other)
    result = run_collect(batch["checklist_path"], other_path)
    # 指纹不匹配 + 该 manifest 里的题号与清单一致，故必然出现 E00
    assert "E00" in codes(result)


def test_collect_reads_row_numbers_and_skips_blank_rows(batch):
    workbook = load_workbook(batch["checklist_path"])
    sheet = workbook[kit.CHECKLIST_SHEET]
    sheet.cell(row=sheet.max_row + 2, column=1, value=None)
    workbook.save(batch["checklist_path"])
    _headers, rows, meta = kit.read_checklist(batch["checklist_path"])
    assert [row[kit.ROW_KEY] for row in rows] == [2, 3, 4]
    assert meta["manifest SHA-256"] == kit.sha256_file(batch["manifest_path"])


# ---------------------------------------------------------------------------
# 回收：最后一道查重
# ---------------------------------------------------------------------------
def test_collect_final_dedup_catches_collision(batch):
    ids = [question_id_at(batch["manifest"], index) for index in range(3)]
    shared_stem = "顺序表的插入与删除操作在最坏情况下分别需要移动多少个元素？"
    edit_rows(
        batch["checklist_path"],
        {
            ids[0]: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_STEM: shared_stem},
            ids[1]: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_STEM: shared_stem},
        },
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert "E20" in codes(result)
    ready = [item["question_id"] for item in result["decisions"] if item["ready_for_apply"]]
    assert ready == [ids[0]]


def test_collect_final_dedup_uses_existing_snapshot(batch):
    qid = question_id_at(batch["manifest"], 0)
    edit_rows(batch["checklist_path"], {qid: {kit.COL_DECISION: kit.DECISION_APPROVE}})
    existing = [{"id": 1, "stem": batch["manifest"]["questions"][0]["stem"]}]
    result = run_collect(batch["checklist_path"], batch["manifest_path"], existing=existing)
    assert "E20" in codes(result)


# ---------------------------------------------------------------------------
# 往返保真
# ---------------------------------------------------------------------------
def test_roundtrip_preserves_unicode_and_newlines(batch):
    qid = question_id_at(batch["manifest"], 0)
    original = batch["manifest"]["questions"][0]
    _headers, rows, _meta = kit.read_checklist(batch["checklist_path"])
    row = next(item for item in rows if kit._norm_cell(item[kit.COL_QID]) == qid)
    assert kit._norm_cell(row[kit.COL_STEM]) == original["stem"]
    assert kit._norm_cell(row[kit.COL_ANALYSIS]) == original["analysis"]
    assert kit._norm_cell(row["选项A"]) == original["options"][0]

    multiline = "第一行结论。\n第二行解释，含全角标点（括号）与 emoji 不允许出现。"
    edit_rows(
        batch["checklist_path"],
        {qid: {kit.COL_DECISION: kit.DECISION_REVISE, kit.COL_STEM: multiline}},
    )
    result = run_collect(batch["checklist_path"], batch["manifest_path"])
    assert result["decisions"][0]["final"]["stem"] == multiline


# ---------------------------------------------------------------------------
# CLI 端到端
# ---------------------------------------------------------------------------
def run_cli(*args, cwd=ROOT):
    return subprocess.run(
        [sys.executable, str(KIT_PATH), *args],
        cwd=str(cwd),
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def test_cli_build_then_collect_end_to_end(batch):
    out_xlsx = batch["tmp_path"] / "cli_checklist.xlsx"
    built = run_cli(
        "build",
        "--manifest", str(batch["manifest_path"]),
        "--out", str(out_xlsx),
        "--built-at", FIXED_TS,
    )
    assert built.returncode == 0, built.stderr
    assert out_xlsx.is_file()

    ids = [question_id_at(batch["manifest"], index) for index in range(3)]
    edit_rows(
        out_xlsx,
        {
            ids[0]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[1]: {kit.COL_DECISION: kit.DECISION_APPROVE},
            ids[2]: {kit.COL_DECISION: kit.DECISION_REJECT},
        },
    )
    out_json = batch["tmp_path"] / "decisions.json"
    collected = run_cli(
        "collect",
        "--checklist", str(out_xlsx),
        "--manifest", str(batch["manifest_path"]),
        "--out", str(out_json),
        "--collected-at", FIXED_TS,
    )
    assert collected.returncode == 0, collected.stderr
    payload = json.loads(out_json.read_text(encoding="utf-8"))
    assert payload["schema_version"] == kit.SCHEMA_VERSION
    assert payload["summary"]["ready_for_apply"] == 2
    assert payload["summary"]["blocking"] == 0


def test_cli_collect_exit_code_one_when_blocked(batch):
    out_json = batch["tmp_path"] / "blocked.json"
    result = run_cli(
        "collect",
        "--checklist", str(batch["checklist_path"]),
        "--manifest", str(batch["manifest_path"]),
        "--out", str(out_json),
    )
    assert result.returncode == 1
    assert out_json.is_file()
    payload = json.loads(out_json.read_text(encoding="utf-8"))
    assert payload["summary"]["blocking"] >= 1


def test_cli_usage_errors_return_two(tmp_path):
    manifest = default_manifest()
    manifest_path = write_manifest(tmp_path, manifest)
    missing = run_cli("build", "--manifest", str(manifest_path))
    assert missing.returncode == 2
    broken = run_cli(
        "collect",
        "--checklist", str(tmp_path / "nope.xlsx"),
        "--manifest", str(manifest_path),
        "--out", str(tmp_path / "x.json"),
    )
    assert broken.returncode == 3


def test_script_never_touches_database_or_app():
    text = KIT_PATH.read_text(encoding="utf-8")
    for forbidden in ("SessionLocal", "create_engine", "alembic"):
        assert forbidden not in text, f"审核工具不应出现 {forbidden}"
    tree = ast.parse(text)
    imported: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module)
        elif isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
    offenders = sorted(name for name in imported if name == "app" or name.startswith("app."))
    assert offenders == [], f"审核工具不应导入后端 app：{offenders}"
