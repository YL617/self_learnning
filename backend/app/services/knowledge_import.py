"""大阶段 4 M2：知识点批量导入引擎。

设计原则（与《M2批量知识库导入设计方案.md》逐条对应）：

* **两条入口、同一条管线**：`preview` 与 `apply` 调用完全相同的
  `parse_* → plan_import`，唯一差别是 `apply` 末尾多一步写入。这保证
  「预览所见 = 执行所得」，也是 TOCTOU 防护的基础（apply 会完整重跑校验）。
* **preview 不落库**：不写任何表、不建批次行、不落盘临时文件。因此 apply
  必须重传同一份数据（服务端刻意不暂存，从根上消除「临时资源生命周期」问题）。
* **全有或全无**：只要存在任何行级 error，整批拒绝、零写入；不做「部分成功」。
* **复用写入逻辑**：本模块**不自己写 INSERT/UPDATE**，知识点的新增与更新一律
  通过 `KnowledgePointService.create()` / `.update()`，避免出现第二套写入口径。
* **零 DB 变更**：不新增任何表/列/约束；批次审计复用 M1 已建好的
  `knowledge_point_import_batches`。

解析层只用标准库（`csv`）+ `openpyxl`（仅 xlsx 分支惰性导入），不引入
pandas / Django 等重型依赖。
"""

from __future__ import annotations

import csv
import io
import re
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import NoReturn

from sqlalchemy import delete as sa_delete
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    KnowledgePoint,
    KnowledgePointImportBatch,
    KnowledgePointPrerequisite,
    QuestionKnowledgePoint,
    UserKnowledgePointMastery,
)
from app.schemas.knowledge import KnowledgePointUpdate
from app.schemas.knowledge_import import (
    BlockingRef,
    PreviewNodeType,
    PreviewRow,
    RowIssue,
)
from app.services.knowledge_point_service import (
    CODE_MAX_LENGTH,
    KP_NODE_TYPE_CONCEPT,
    KP_NODE_TYPE_CONTAINER,
    KP_SOURCE_ADMIN,
    KnowledgePointService,
    clean_name,
    clean_subject,
    normalize_subject,
)

# --------------------------------------------------------------------------
# 上限与常量
# --------------------------------------------------------------------------

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 500
PREVIEW_ROW_LIMIT = 200

SUBJECT_MAX_LENGTH = 100
NAME_MAX_LENGTH = 200
MINUTES_MAX = 100_000
ALIAS_MAX_ITEMS = 20
ALIAS_MAX_LENGTH = 100
MAX_DEPTH = 12
DEEP_HIERARCHY_THRESHOLD = 4

STRATEGY_SKIP = "skip"
STRATEGY_UPDATE_EMPTY = "update_empty"
CONFLICT_STRATEGIES: tuple[str, ...] = (STRATEGY_SKIP, STRATEGY_UPDATE_EMPTY)

# `.xls` 是旧二进制格式，与 xlsx（zip）完全不同，必须显式拒绝而不是「解析失败」。
LEGACY_XLS_SUFFIX = ".xls"
ALLOWED_IMPORT_SUFFIXES: tuple[str, ...] = (".csv", ".tsv", ".txt", ".xlsx")

# 无表头时的列位置（顺序即位置）。`code` **追加在末位**是刻意的兼容性决策：
# 早期文档已把 1~6 位定为「学科/路径/名/别名/难度/学时」，插在中间会让按位置
# 粘贴的历史习惯整体错位。
COLUMN_SPEC: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("subject", ("学科", "科目", "subject")),
    ("parent_path", ("层级路径", "路径", "层级", "parent_path", "path")),
    ("name", ("知识点名", "知识点名称", "名称", "name")),
    ("aliases", ("别名", "alias", "aliases")),
    ("difficulty", ("难度", "difficulty")),
    ("estimated_minutes", ("预计学时", "预估时长", "时长", "estimated_minutes")),
    ("code", ("编码", "code")),
)
CANONICAL_FIELDS: tuple[str, ...] = tuple(name for name, _ in COLUMN_SPEC)

_HEADER_ALIAS_TO_FIELD: dict[str, str] = {
    alias.lower(): field_name for field_name, aliases in COLUMN_SPEC for alias in aliases
}

# 难度中文同义词 → 词表值。词表本身仍由 KP_DIFFICULTY_LEVELS 与 DB CHECK 钉死，
# 这里只做「输入友好化」，不引入第二套难度语义。
DIFFICULTY_SYNONYMS: dict[str, str] = {
    "easy": "easy",
    "简单": "easy",
    "容易": "easy",
    "medium": "medium",
    "中等": "medium",
    "一般": "medium",
    "hard": "hard",
    "困难": "hard",
    "难": "hard",
}

DEFAULT_DELIMITERS: dict[str, str] = {".tsv": "\t", ".txt": "\t", ".csv": ","}
DELIMITER_LABELS = {"\t": "制表符", ",": "逗号", ";": "分号", "|": "竖线"}

PATH_SEPARATOR_RE = re.compile(r"[>＞]")
ALIAS_SEPARATOR_RE = re.compile(r"[|,，]")
INTEGER_RE = re.compile(r"\d+(?:\.0+)?")

# 节点类型推导来源文案（大阶段 4 P0）：让管理员在 preview 里看懂「为什么是这个类型」。
NODE_TYPE_SOURCE_ROW = "知识点行"
NODE_TYPE_SOURCE_PATH = "由层级路径自动推导"


# --------------------------------------------------------------------------
# 异常
# --------------------------------------------------------------------------


class KnowledgeImportError(Exception):
    """导入管线可预期失败的基类（载体：结构化 code + 明细）。"""

    code = "B00"

    def __init__(
        self,
        message: str,
        *,
        code: str | None = None,
        issues: list[RowIssue] | None = None,
        cycle: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.message = message
        if code is not None:
            self.code = code
        self.issues: list[RowIssue] = list(issues or [])
        self.cycle: list[str] = list(cycle or [])


class ImportParseError(KnowledgeImportError):
    """B01：编码不可识别 / 不是 xlsx / 旧版 .xls / 空文件 / 二进制噪声。"""

    code = "B01"


class ImportStructureError(KnowledgeImportError):
    """B02 / B03：表头结构非法、行数超限。"""


class ImportCycleError(KnowledgeImportError):
    """B04：整批层级成环。"""

    code = "B04"


class ImportConflictError(KnowledgeImportError):
    """apply 阶段的可预期冲突（行级 error / 行数不匹配 / 回滚受阻）。"""

    code = "A01"


class ImportWriteError(KnowledgeImportError):
    """写入期异常（已整体回滚，批次以 failed 落库）。"""

    code = "A02"


class ImportBatchNotFound(KnowledgeImportError):
    code = "A03"


# --------------------------------------------------------------------------
# 数据载体
# --------------------------------------------------------------------------


@dataclass
class ParsedRow:
    """一行数据。`cells` 是原始单元格，其余字段由校验阶段填充。"""

    row: int
    cells: dict[str, str] = field(default_factory=dict)
    column_overflow: int = 0

    subject: str = ""
    normalized_subject: str = ""
    parent_path: str | None = None
    # 本行真正的「父级候选链」：即去掉「末段 == 本行名」之后的路径段。
    parent_chain: list[str] = field(default_factory=list)
    name: str = ""
    normalized_name: str = ""
    aliases: list[str] = field(default_factory=list)
    difficulty: str | None = None
    estimated_minutes: int | None = None
    code: str | None = None
    # 行级节点类型：本行的「知识点名」一律是 concept（目录由层级路径推导为 container）。
    node_type: str = KP_NODE_TYPE_CONCEPT
    node_type_source: str = NODE_TYPE_SOURCE_ROW

    # 行内纯校验结论
    issues: list[RowIssue] = field(default_factory=list)
    # 查库后才能得出的结论（E08 / W01~W03），赋值而非追加，保证 plan 可重复调用
    db_issues: list[RowIssue] = field(default_factory=list)

    @property
    def all_issues(self) -> list[RowIssue]:
        return [*self.issues, *self.db_issues]

    @property
    def has_error(self) -> bool:
        return any(issue.level == "error" for issue in self.all_issues)


@dataclass
class ParsedTable:
    source_format: str
    detected_encoding: str | None
    detected_delimiter: str | None
    has_header: bool
    rows: list[ParsedRow] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# 解析层（纯函数，可脱库单测）
# --------------------------------------------------------------------------


def _collapse(value: str) -> str:
    """复用知识点的空白清洗口径（全角空格→半角、连续空白折叠、去首尾）。

    `clean_subject` 就是这个原语；直接复用它，避免在导入侧再造一套清洗规则。
    """
    return clean_subject(value)


def _display(raw: str | None) -> str:
    return _collapse(raw or "")


def _normalize_header(cell: str) -> str:
    # 兜底剥离 BOM：`utf-8-sig` 已处理，但粘贴路径可能把 BOM 带进来。
    return cell.lstrip("\ufeff").strip().lower()


def _looks_like_header(cells: list[str]) -> bool:
    """表头判据：至少 2 个单元格命中已知列名。

    要求 ≥2 命中而不是「第一格命中」，是为了降低误判 —— 纯数据行几乎不可能
    连续出现两个与列名同名的值。
    """
    hits = sum(1 for cell in cells if _normalize_header(cell) in _HEADER_ALIAS_TO_FIELD)
    return hits >= 2


def _decode(raw: bytes) -> tuple[str, str, str | None]:
    """返回 `(文本, 实际使用的编码, 提示)`。

    顺序**不可颠倒**：`gb18030` 能「成功」解码几乎任意字节序列，放在前面会把
    真正的 UTF-8 内容解成乱码却不报错（R17）。
    """
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "utf-8-sig", None
    try:
        return raw.decode("utf-8"), "utf-8", None
    except UnicodeDecodeError:
        pass
    try:
        text = raw.decode("gb18030")
    except UnicodeDecodeError as exc:
        raise ImportParseError(
            "无法识别文件编码，请另存为 UTF-8 或 CSV UTF-8 后重试"
        ) from exc
    return text, "gb18030", "编码按 gb18030 解析（疑似 GBK/ANSI），请核对中文是否正常"


def _detect_delimiter(text: str, *, suffix: str) -> str:
    """按「内容」判定分隔符，扩展名只作平票时的倾向。

    刻意不用 `csv.Sniffer`：它在只有两行、且各行字段数不一致的小样本上会给出
    离谱结论（实测把 `学科,知识点名\\n数据结构,栈,x,y,z` 判成制表符分隔）。这里
    改成对 4 个候选分隔符逐一试切，再按
    「列数一致性（众数出现次数）→ 众数列数 → 是否等于扩展名默认值」
    排序取最优 —— 判定可解释，且不会被「某一列里恰好大量出现 `|`」这种数据骗到。
    """
    sample_lines = [line for line in text.splitlines()[:50] if line.strip()]
    fallback = DEFAULT_DELIMITERS.get(suffix, "\t")
    best: tuple[tuple[int, int, int], str] | None = None
    for candidate in (",", "\t", ";", "|"):
        counts = [len(next(csv.reader([line], delimiter=candidate))) for line in sample_lines]
        tallies: dict[int, int] = {}
        for count in counts:
            tallies[count] = tallies.get(count, 0) + 1
        if not tallies:
            continue
        modal_count, modal_hits = max(tallies.items(), key=lambda item: (item[1], item[0]))
        if modal_count < 2:
            continue
        score = (modal_hits, modal_count, 1 if candidate == fallback else 0)
        if best is None or score > best[0]:
            best = (score, candidate)
    return best[1] if best else fallback


def _delimiter_label(delimiter: str) -> str:
    return DELIMITER_LABELS.get(delimiter, delimiter)


def _xlsx_cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        # 让 Excel 里的 30.0 还原成 "30"，避免后续整数解析被误判为非法。
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value)


def _parse_xlsx(raw: bytes, *, notes: list[str]) -> list[list[str]]:
    """xlsx 分支：仅此处依赖 openpyxl（惰性导入，避免拖垮应用启动）。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失时的兜底
        raise ImportParseError(
            "服务端缺少 xlsx 解析依赖（openpyxl），请改用 CSV/TSV 上传"
        ) from exc

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 的失败种类很多，统一转成 400
        message = str(exc) or exc.__class__.__name__
        if "zip" in message.lower():
            raise ImportParseError(
                "文件不是有效的 xlsx（可能已损坏，或只是把别的文件改了后缀名）"
            ) from exc
        raise ImportParseError(f"无法解析 xlsx：{message}") from exc

    try:
        if len(workbook.sheetnames) > 1:
            notes.append(
                f"工作簿含 {len(workbook.sheetnames)} 个工作表，只读取第一个"
                f"「{workbook.sheetnames[0]}」"
            )
        sheet = workbook[workbook.sheetnames[0]]
        matrix: list[list[str]] = []
        # read_only 下 ws.max_row 可能被虚报（声明范围），因此边迭代边计数；
        # 并且必须在 finally 里 close()，否则句柄泄漏直到进程耗尽文件描述符。
        for raw_row in sheet.iter_rows(values_only=True):
            matrix.append([_xlsx_cell_to_text(value) for value in raw_row])
    finally:
        workbook.close()
    return matrix


def _parse_delimited(text: str, *, suffix: str, notes: list[str]) -> tuple[list[list[str]], str]:
    if "\x00" in text:
        raise ImportParseError("文件内容不是可解析的文本（检测到二进制内容）")
    delimiter = _detect_delimiter(text, suffix=suffix)
    notes.append(f"分隔符按「{_delimiter_label(delimiter)}」解析")
    expected = DEFAULT_DELIMITERS.get(suffix)
    if expected is not None and expected != delimiter:
        notes.append(
            f"内容分隔符（{_delimiter_label(delimiter)}）与扩展名 {suffix} 不一致，已按内容解析"
        )
    # 一律用标准库 csv.reader（RFC 4180：引号、转义引号、字段内换行都由它处理），
    # 绝不自己 split —— 生产者漏写引号导致的「行错位」会在列数上暴露出来。
    matrix = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return matrix, delimiter


def _finish_table(
    matrix: list[list[str]],
    *,
    source_format: str,
    detected_encoding: str | None,
    detected_delimiter: str | None,
    notes: list[str],
) -> ParsedTable:
    """规范成 `ParsedRow` 列表：定表头 → 定列 → 剔空行 → 行数上限。"""
    indexed = [(index, row) for index, row in enumerate(matrix, start=1)]
    kept = [(index, row) for index, row in indexed if any(cell.strip() for cell in row)]
    if not kept:
        raise ImportParseError("文件中没有任何数据行")
    blank_rows = len(indexed) - len(kept)
    if blank_rows:
        notes.append(f"已剔除 {blank_rows} 个空行")

    _first_index, first_cells = kept[0]
    has_header = _looks_like_header(first_cells)
    data_rows = kept[1:] if has_header else kept

    field_index: dict[str, int] = {}
    if has_header:

        def _header_error(message: str) -> NoReturn:
            # 整批级结论：`row=None` 让调用方把它归入 batch_errors 而不是行级问题。
            raise ImportStructureError(
                message,
                code="B02",
                issues=[RowIssue(row=None, level="error", code="B02", message=message)],
            )

        for position, cell in enumerate(first_cells):
            field_name = _HEADER_ALIAS_TO_FIELD.get(_normalize_header(cell))
            if field_name is None:
                continue
            if field_name in field_index:
                _header_error(f"表头存在重复列「{cell.strip()}」，请修改后重试")
            field_index[field_name] = position
        missing = [
            label
            for label, field_name in (("学科", "subject"), ("知识点名", "name"))
            if field_name not in field_index
        ]
        if missing:
            _header_error(f"表头缺少必需列：{'、'.join(missing)}")
        if len(first_cells) < 2:
            _header_error("表头至少需要 2 列（学科、知识点名）")
        column_count = len(first_cells)
    else:
        field_index = {name: position for position, name in enumerate(CANONICAL_FIELDS)}
        column_count = max((len(row) for _, row in data_rows), default=0)
        notes.append(
            "未检测到表头，已按列位置解析"
            "（学科 / 层级路径 / 知识点名 / 别名 / 难度 / 学时 / 编码）"
        )

    rows: list[ParsedRow] = []
    for index, cells in data_rows:
        values = {
            name: (cells[position] if position < len(cells) else "")
            for name, position in field_index.items()
        }
        rows.append(
            ParsedRow(
                row=index,
                cells=values,
                column_overflow=max(0, len(cells) - column_count),
            )
        )
    if not rows:
        raise ImportParseError("文件中没有任何数据行")
    if len(rows) > MAX_IMPORT_ROWS:
        message = f"数据行数 {len(rows)} 超过单批上限 {MAX_IMPORT_ROWS}"
        raise ImportStructureError(
            f"单批最多导入 {MAX_IMPORT_ROWS} 行，本次解析出 {len(rows)} 行，请拆分后重试",
            code="B03",
            issues=[RowIssue(row=None, level="error", code="B03", message=message)],
        )
    return ParsedTable(
        source_format=source_format,
        detected_encoding=detected_encoding,
        detected_delimiter=detected_delimiter,
        has_header=has_header,
        rows=rows,
        notes=notes,
    )


def parse_pasted(text: str) -> ParsedTable:
    """入口 A：粘贴的文本。"""
    if not text.strip():
        raise ImportParseError("粘贴内容为空")
    notes: list[str] = []
    matrix, delimiter = _parse_delimited(text, suffix="", notes=notes)
    return _finish_table(
        matrix,
        source_format="paste",
        detected_encoding=None,
        detected_delimiter=delimiter,
        notes=notes,
    )


def parse_uploaded(raw: bytes, filename: str | None) -> ParsedTable:
    """入口 B：上传的文件（.csv / .tsv / .txt / .xlsx）。"""
    suffix = ""
    if filename and "." in filename:
        suffix = "." + filename.rsplit(".", 1)[-1].lower()
    if suffix == LEGACY_XLS_SUFFIX:
        raise ImportParseError(
            "不支持旧版 .xls 格式，请在 Excel 中另存为「CSV UTF-8」或 .xlsx 后重试"
        )
    if suffix not in ALLOWED_IMPORT_SUFFIXES:
        raise ImportParseError(
            "仅支持 .csv / .tsv / .txt / .xlsx，请在 Excel 中另存为「CSV UTF-8」后重试"
        )
    if not raw:
        raise ImportParseError("上传的文件为空")

    notes: list[str] = []
    if suffix == ".xlsx":
        matrix = _parse_xlsx(raw, notes=notes)
        return _finish_table(
            matrix,
            source_format="xlsx",
            detected_encoding=None,
            detected_delimiter=None,
            notes=notes,
        )
    text, encoding, note = _decode(raw)
    if note:
        notes.append(note)
    matrix, delimiter = _parse_delimited(text, suffix=suffix, notes=notes)
    return _finish_table(
        matrix,
        source_format=suffix.lstrip("."),
        detected_encoding=encoding,
        detected_delimiter=delimiter,
        notes=notes,
    )


# --------------------------------------------------------------------------
# 行级校验（纯函数）
# --------------------------------------------------------------------------


def _split_path(raw: str) -> list[str]:
    return [
        segment
        for segment in (_collapse(part) for part in PATH_SEPARATOR_RE.split(raw))
        if segment
    ]


def _split_aliases(raw: str) -> list[str]:
    return [item for item in (_collapse(part) for part in ALIAS_SEPARATOR_RE.split(raw)) if item]


def _parse_difficulty(raw: str) -> tuple[str | None, str | None]:
    if not raw:
        return None, None
    value = DIFFICULTY_SYNONYMS.get(raw.strip().lower())
    if value is None:
        return None, "难度只能是 easy / medium / hard（或 简单 / 中等 / 困难）"
    return value, None


def _parse_minutes(raw: str) -> tuple[int | None, str | None]:
    if not raw:
        return None, None
    text = raw.strip()
    if not INTEGER_RE.fullmatch(text):
        return None, "预计学时必须是不小于 0 的整数分钟"
    value = int(float(text))
    if value > MINUTES_MAX:
        return None, f"预计学时不能超过 {MINUTES_MAX} 分钟"
    return value, None


def validate_rows(rows: list[ParsedRow]) -> None:
    """行级校验（E01–E07、E09–E11、W04–W06），就地写回清洗后的字段。

    幂等：每次调用先清空 `issues` / `db_issues`，因此 preview 与 apply 可以
    各跑一次而不产生重复结论。
    """
    for row in rows:
        row.issues = []
        row.db_issues = []

    seen: dict[tuple[str, str], int] = {}
    for row in rows:
        cells = row.cells
        row.subject = _display(cells.get("subject", ""))
        row.normalized_subject = normalize_subject(row.subject) if row.subject else ""
        row.name = _display(cells.get("name", ""))
        row.normalized_name = clean_name(row.name) if row.name else ""
        row.parent_path = _display(cells.get("parent_path", "")) or None

        if row.column_overflow > 0:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E11",
                    message=(
                        f"该行比表头多出 {row.column_overflow} 列，"
                        "疑似分隔符或引号导致的列错位"
                    ),
                )
            )
        if not row.subject:
            row.issues.append(
                RowIssue(row=row.row, level="error", code="E01", field="subject", message="学科不能为空")
            )
        elif len(row.subject) > SUBJECT_MAX_LENGTH:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E03",
                    field="subject",
                    message=f"学科超过 {SUBJECT_MAX_LENGTH} 个字符",
                )
            )
        if not row.name:
            row.issues.append(
                RowIssue(row=row.row, level="error", code="E02", field="name", message="知识点名称不能为空")
            )
        elif len(row.name) > NAME_MAX_LENGTH:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E03",
                    field="name",
                    message=f"知识点名称超过 {NAME_MAX_LENGTH} 个字符",
                )
            )

        row.code = _collapse(cells.get("code", "")) or None
        if row.code and len(row.code) > CODE_MAX_LENGTH:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E04",
                    field="code",
                    message=f"编码超过 {CODE_MAX_LENGTH} 个字符",
                )
            )
            row.code = None

        row.difficulty, difficulty_error = _parse_difficulty(cells.get("difficulty", ""))
        if difficulty_error:
            row.issues.append(
                RowIssue(
                    row=row.row, level="error", code="E05", field="difficulty", message=difficulty_error
                )
            )
        row.estimated_minutes, minutes_error = _parse_minutes(
            cells.get("estimated_minutes", "")
        )
        if minutes_error:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E06",
                    field="estimated_minutes",
                    message=minutes_error,
                )
            )

        row.aliases = _split_aliases(cells.get("aliases", ""))
        if len(row.aliases) > ALIAS_MAX_ITEMS or any(
            len(alias) > ALIAS_MAX_LENGTH for alias in row.aliases
        ):
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E07",
                    field="aliases",
                    message=f"别名最多 {ALIAS_MAX_ITEMS} 个，单个不超过 {ALIAS_MAX_LENGTH} 个字符",
                )
            )

        # 层级路径：末段与本行同名表示「本行就是路径末段」，父级取倒数第二段。
        segments = _split_path(cells.get("parent_path", ""))
        if segments and clean_name(segments[-1]) == row.normalized_name:
            chain = segments[:-1]
        else:
            chain = segments
        row.parent_chain = chain
        if len(segments) > MAX_DEPTH:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E03",
                    field="parent_path",
                    message=f"层级路径超过 {MAX_DEPTH} 层，请简化后重试",
                )
            )
        if any(clean_name(segment) == row.normalized_name for segment in chain):
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E10",
                    field="parent_path",
                    message="层级路径不能把知识点自己作为父级",
                )
            )
        if len(segments) > DEEP_HIERARCHY_THRESHOLD:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="warning",
                    code="W04",
                    field="parent_path",
                    message=(
                        f"层级深度 {len(segments)} 超过建议值 {DEEP_HIERARCHY_THRESHOLD}，"
                        "不利于推荐与展示"
                    ),
                )
            )

        # E09：批内 (学科, 名称) 重复 —— 保留首次出现，后续重复行报错并给出首行号。
        if row.normalized_subject and row.normalized_name:
            key = (row.normalized_subject, row.normalized_name)
            first_row = seen.get(key)
            if first_row is None:
                seen[key] = row.row
            else:
                row.issues.append(
                    RowIssue(
                        row=row.row,
                        level="error",
                        code="E09",
                        field="name",
                        message=f"与第 {first_row} 行重复（同一学科下同名知识点）",
                    )
                )

        if row.difficulty is None and row.estimated_minutes is None:
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="warning",
                    code="W06",
                    field="difficulty",
                    message="未提供难度与预计学时，推荐与展示效果会退化",
                )
            )

    _mark_alias_cross_row(rows)
    _mark_node_type_conflicts(rows)


def _mark_node_type_conflicts(rows: list[ParsedRow]) -> None:
    """E12（I3）：同一名称在本批内既被作为「目录」（路径段）又被作为「知识点」声明。

    例：
        行1：路径「栈」，知识点「括号匹配」   → 栈 = container
        行2：路径空，  知识点「栈」           → 栈 = concept
    这属于**节点身份冲突**，不做静默裁决（不"显式行优先"），而是整批报 error，
    拒绝 apply，强制管理员先人工统一知识模型。这是「零污染」的硬门禁。
    """
    concept_rows: dict[str, list[int]] = {}
    container_rows: dict[str, list[int]] = {}
    for row in rows:
        if row.normalized_name:
            concept_rows.setdefault(row.normalized_name, []).append(row.row)
        for segment in row.parent_chain:
            key = clean_name(segment)
            if key:
                container_rows.setdefault(key, []).append(row.row)

    conflicts = set(concept_rows) & set(container_rows)
    if not conflicts:
        return

    for row in rows:
        declared: set[str] = set()
        if row.normalized_name in conflicts:
            declared.add(row.normalized_name)
            others = "、".join(
                f"第 {item} 行" for item in sorted(set(container_rows[row.normalized_name]))
            )
            row.issues.append(
                RowIssue(
                    row=row.row,
                    level="error",
                    code="E12",
                    field="name",
                    message=(
                        f"「{row.name}」在本批中同时被作为目录（{others}的层级路径）"
                        "与知识点声明，节点类型冲突，请先人工统一知识模型"
                    ),
                )
            )
        for segment in row.parent_chain:
            key = clean_name(segment)
            if key in conflicts and key not in declared:
                declared.add(key)
                others = "、".join(
                    f"第 {item} 行" for item in sorted(set(concept_rows[key]))
                )
                row.issues.append(
                    RowIssue(
                        row=row.row,
                        level="error",
                        code="E12",
                        field="parent_path",
                        message=(
                            f"层级路径中的「{segment}」在本批中同时被作为知识点声明"
                            f"（{others}），节点类型冲突，请先人工统一知识模型"
                        ),
                    )
                )


def _mark_alias_cross_row(rows: list[ParsedRow]) -> None:
    """W05：跨行别名冲突（M1 有意留到 M2 的检查；只提示，不自动合并）。"""
    index: dict[str, set[int]] = {}
    for row in rows:
        if row.normalized_name:
            index.setdefault(row.normalized_name, set()).add(row.row)
        for alias in row.aliases:
            index.setdefault(clean_name(alias), set()).add(row.row)
    for row in rows:
        for alias in row.aliases:
            others = sorted(index.get(clean_name(alias), set()) - {row.row})
            if others:
                rows_text = "、".join(f"第 {item} 行" for item in others)
                row.issues.append(
                    RowIssue(
                        row=row.row,
                        level="warning",
                        code="W05",
                        field="aliases",
                        message=(
                            f"别名「{alias}」与{rows_text}的名称或别名重复，"
                            "建议合并为同一条知识点"
                        ),
                    )
                )
                break


# --------------------------------------------------------------------------
# 计划层（只读查库，preview 与 apply 共用）
# --------------------------------------------------------------------------


@dataclass
class PlannedRow:
    parsed: ParsedRow
    action: str
    existing: KnowledgePoint | None = None
    parent_node: tuple | None = None


@dataclass
class AutoParent:
    key: tuple[str, str, str]
    parent_key: tuple | None
    subject: str
    name: str
    display_path: str
    depth: int
    order: int


@dataclass
class PlanResult:
    table: ParsedTable
    rows: list[PlannedRow] = field(default_factory=list)
    auto_parents: list[AutoParent] = field(default_factory=list)
    new_subjects: list[str] = field(default_factory=list)
    cycle: list[str] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)

    @property
    def total_rows(self) -> int:
        return len(self.rows)

    @property
    def error_rows(self) -> int:
        return sum(1 for item in self.rows if item.parsed.has_error)

    @property
    def warning_rows(self) -> int:
        return sum(
            1
            for item in self.rows
            if any(issue.level == "warning" for issue in item.parsed.all_issues)
        )


def _kp_node(knowledge_point_id: int) -> tuple[str, int]:
    return ("kp", knowledge_point_id)


def _new_node(normalized_subject: str, normalized_name: str) -> tuple[str, str, str]:
    return ("new", normalized_subject, normalized_name)


def plan_import(db: Session, table: ParsedTable, strategy: str) -> PlanResult:
    """查库生成执行计划。**只读**，preview 与 apply 完全共用。"""
    if strategy not in CONFLICT_STRATEGIES:
        raise ImportStructureError(f"未知的冲突策略：{strategy}", code="B02")

    validate_rows(table.rows)
    plan = PlanResult(table=table)

    names: set[str] = set()
    subjects: set[str] = set()
    for row in table.rows:
        if row.normalized_name:
            names.add(row.normalized_name)
        names.update(clean_name(segment) for segment in row.parent_chain)
        if row.normalized_subject:
            subjects.add(row.normalized_subject)

    existing_by_key: dict[tuple[str, str], KnowledgePoint] = {}
    by_name: dict[str, list[KnowledgePoint]] = {}
    statement = select(KnowledgePoint).where(
        KnowledgePoint.normalized_subject.in_(subjects or {""})
        | KnowledgePoint.normalized_name.in_(names or {""})
    )
    for item in db.scalars(statement).all():
        existing_by_key[(item.normalized_subject, item.normalized_name)] = item
        by_name.setdefault(item.normalized_name, []).append(item)

    # W01 的判据：该学科是否已有知识点（按学科探测，不做全表扫描）。
    for normalized_subject in sorted(subjects):
        exists = db.scalar(
            select(KnowledgePoint.id)
            .where(KnowledgePoint.normalized_subject == normalized_subject)
            .limit(1)
        )
        if exists is None:
            display = next(
                (row.subject for row in table.rows if row.normalized_subject == normalized_subject),
                normalized_subject,
            )
            plan.new_subjects.append(display)

    # ---- 第一遍：判定每行动作 ----
    for row in table.rows:
        existing = (
            existing_by_key.get((row.normalized_subject, row.normalized_name))
            if row.normalized_subject and row.normalized_name
            else None
        )
        if existing is None:
            action = "create"
        elif strategy == STRATEGY_UPDATE_EMPTY:
            action = "update_empty"
        else:
            action = "skip"
        plan.rows.append(PlannedRow(parsed=row, action=action, existing=existing))
        if existing is not None:
            if existing.node_type != KP_NODE_TYPE_CONCEPT:
                # I4：已有节点不得因本次导入改变类型。库中的目录节点被本批当作
                # 知识点（concept）导入 = **节点身份冲突** → 整批 error（不是 warning，
                # 不允许"确认后硬导"，必须先人工解决知识模型）。
                row.db_issues = [
                    RowIssue(
                        row=row.row,
                        level="error",
                        code="E13",
                        field="name",
                        message=(
                            f"「{existing.name}」在库中已是目录节点，"
                            "本批把它作为知识点导入会造成节点类型冲突，"
                            "请先人工解决知识模型"
                        ),
                    )
                ]
            else:
                row.db_issues = [
                    RowIssue(
                        row=row.row,
                        level="warning",
                        code="W03",
                        field="name",
                        message=(
                            f"已存在同名知识点（id={existing.id}），"
                            + ("将补充其空字段" if action == "update_empty" else "本次将跳过")
                        ),
                    )
                ]

    # ---- 节点索引：既有知识点 → 本批新建行 → 自动创建父节点（逐层追加）----
    node_for: dict[tuple[str, str], tuple] = {
        key: _kp_node(item.id) for key, item in existing_by_key.items()
    }
    labels: dict[tuple, str] = {
        key: f"{item.subject} > {item.name}" for key, item in existing_by_key.items()
    }
    for planned in plan.rows:
        row = planned.parsed
        if planned.action == "create" and not row.has_error:
            key = (row.normalized_subject, row.normalized_name)
            node_for.setdefault(key, _new_node(*key))
            labels.setdefault(_new_node(*key), f"{row.subject} > {row.name}")

    # ---- 第二遍：逐层解析父级链 ----
    auto_by_key: dict[tuple, AutoParent] = {}
    batch_parent: dict[tuple, tuple | None] = {}
    for planned in plan.rows:
        row = planned.parsed
        if row.has_error:
            continue
        # 不写库的行（skip / update_empty）不新建父节点 —— 否则会留下没有任何
        # 行引用的孤儿节点；同时继续走完父级链，只为把 E08 查清。
        # 这样 preview 的 create_parent 与 apply 的真实行为始终一致。
        will_create = planned.action == "create"
        resolved: tuple | None = None
        auto_created: list[str] = []
        for depth, segment in enumerate(row.parent_chain, start=1):
            segment_key = (row.normalized_subject, clean_name(segment))
            if segment_key in node_for:
                resolved = node_for[segment_key]
                continue
            cross = [
                item
                for item in by_name.get(segment_key[1], [])
                if item.normalized_subject != row.normalized_subject
            ]
            if cross:
                subjects_text = "、".join(sorted({item.subject for item in cross}))
                row.db_issues.append(
                    RowIssue(
                        row=row.row,
                        level="error",
                        code="E08",
                        field="parent_path",
                        message=(
                            f"父知识点「{segment}」属于学科「{subjects_text}」，"
                            f"与「{row.subject}」不一致"
                        ),
                    )
                )
                resolved = None
                break
            if not will_create:
                resolved = None
                continue
            auto_key = _new_node(*segment_key)
            auto = auto_by_key.get(auto_key)
            if auto is None:
                path = " > ".join(_collapse(part) for part in row.parent_chain[:depth])
                auto = AutoParent(
                    key=auto_key,
                    parent_key=resolved,
                    subject=row.subject,
                    name=_collapse(segment),
                    display_path=path,
                    depth=depth,
                    order=len(plan.auto_parents),
                )
                plan.auto_parents.append(auto)
                auto_by_key[auto_key] = auto
                labels.setdefault(auto_key, f"{row.subject} > {_collapse(segment)}")
            auto_created.append(auto.display_path)
            node_for[segment_key] = auto_key
            resolved = auto_key
        planned.parent_node = resolved
        if planned.action == "create" and not row.has_error:
            batch_parent[_new_node(row.normalized_subject, row.normalized_name)] = resolved
        if auto_created:
            row.db_issues.append(
                RowIssue(
                    row=row.row,
                    level="warning",
                    code="W02",
                    field="parent_path",
                    message="将自动创建父节点：" + "、".join(dict.fromkeys(auto_created)),
                )
            )

    # W01：学科在库中不存在。
    new_subject_keys = {normalize_subject(item) for item in plan.new_subjects}
    for planned in plan.rows:
        row = planned.parsed
        if row.normalized_subject and row.normalized_subject in new_subject_keys:
            row.db_issues.append(
                RowIssue(
                    row=row.row,
                    level="warning",
                    code="W01",
                    field="subject",
                    message=f"学科「{row.subject}」在库中尚不存在，本次将新建",
                )
            )

    plan.auto_parents.sort(key=lambda item: (item.depth, item.order))
    plan.cycle = _detect_cycle(batch_parent, labels) or []
    plan.counts = {
        "create": sum(1 for item in plan.rows if item.action == "create"),
        "skip": sum(1 for item in plan.rows if item.action == "skip"),
        "update_empty": sum(1 for item in plan.rows if item.action == "update_empty"),
        "create_parent": len(plan.auto_parents),
        # 显式的双计数：知识点（concept）与目录（container）分开，避免混成一个 created_count。
        "create_concept": sum(1 for item in plan.rows if item.action == "create"),
        "create_container": len(plan.auto_parents),
    }
    return plan


def _detect_cycle(
    batch_parent: dict[tuple, tuple | None], labels: dict[tuple, str]
) -> list[str] | None:
    """在「本批新建节点」的子图上找环并回溯出环路径。

    既有知识点之间不可能成环（`_assert_no_cycle` + DB 约束保证），且本批无法
    修改既有节点的父级，因此环必然完全落在本批新建节点内 —— 只在这个子图上
    DFS 即可，无需全表扫描。
    """
    for start in batch_parent:
        path: list[tuple] = []
        index: dict[tuple, int] = {}
        current: tuple | None = start
        while current is not None and current in batch_parent:
            if current in index:
                cycle = path[index[current] :]
                return [labels.get(node, "?") for node in cycle] + [
                    labels.get(cycle[0], "?")
                ]
            index[current] = len(path)
            path.append(current)
            current = batch_parent[current]
    return None


def preview_row(planned: PlannedRow) -> PreviewRow:
    row = planned.parsed
    return PreviewRow(
        row=row.row,
        subject=row.subject,
        parent_path=row.parent_path,
        name=row.name,
        code=row.code,
        aliases=list(row.aliases),
        difficulty=row.difficulty,
        estimated_minutes=row.estimated_minutes,
        node_type=row.node_type,
        node_type_source=row.node_type_source,
        action=planned.action,  # type: ignore[arg-type]
        existing_kp_id=planned.existing.id if planned.existing is not None else None,
        issues=row.all_issues,
    )


def preview_rows(plan: PlanResult, *, limit: int = PREVIEW_ROW_LIMIT) -> list[PreviewRow]:
    return [preview_row(item) for item in plan.rows[:limit]]


def preview_auto_parent_nodes(plan: PlanResult) -> list[PreviewNodeType]:
    """预览里显式列出「将自动创建的目录节点」及其类型来源。"""
    return [
        PreviewNodeType(
            name=auto.name,
            path=auto.display_path,
            node_type=KP_NODE_TYPE_CONTAINER,
            node_type_source=NODE_TYPE_SOURCE_PATH,
        )
        for auto in plan.auto_parents
    ]


# --------------------------------------------------------------------------
# 写入层（apply）
# --------------------------------------------------------------------------


@dataclass
class ApplyReport:
    batch_id: int
    total_rows: int
    created_count: int
    updated_count: int
    skipped_count: int
    auto_parent_count: int
    duration_ms: int
    rows: list[PreviewRow] = field(default_factory=list)


@dataclass
class _CreationSpec:
    """一个「待创建的节点」：可能是自动父节点，也可能是本批的目标知识点。"""

    subject: str
    name: str
    parent_key: tuple | None
    extras: dict[str, object] = field(default_factory=dict)


def _creation_specs(plan: PlanResult) -> dict[tuple, _CreationSpec]:
    """把「自动父节点 + 待新建的行」统一成一张 id 待定表。

    统一成一张表的原因：父级既可能是既有知识点、也可能是**同批次里另一行**
    要建的知识点。若分两趟先建父再建行，就会出现「父节点其实是本批的某一行、
    此时还没有 id」的空洞（这正是首版实现被测试抓出来的缺陷）。
    """
    specs: dict[tuple, _CreationSpec] = {}
    for auto in plan.auto_parents:
        # 自动创建的父节点 = 纯目录节点（container），只参与层级与展示。
        specs[auto.key] = _CreationSpec(
            subject=auto.subject,
            name=auto.name,
            parent_key=auto.parent_key,
            extras={"node_type": KP_NODE_TYPE_CONTAINER},
        )
    for planned in plan.rows:
        if planned.action != "create":
            continue
        row = planned.parsed
        specs[_new_node(row.normalized_subject, row.normalized_name)] = _CreationSpec(
            subject=row.subject,
            name=row.name,
            parent_key=planned.parent_node,
            extras={
                # 真实导入行 = 可学习知识点（concept）。
                "node_type": KP_NODE_TYPE_CONCEPT,
                "code": row.code,
                "aliases": row.aliases or None,
                "difficulty": row.difficulty,
                "estimated_minutes": row.estimated_minutes,
            },
        )
    return specs


def apply_import(
    db: Session,
    *,
    user_id: int,
    table: ParsedTable,
    strategy: str,
    source_name: str | None,
    expected_total_rows: int | None,
) -> ApplyReport:
    """重新解析/校验/计划 → 单事务写入 → 批次落库。任一行 error 则整批拒绝。"""
    started = time.monotonic()
    plan = plan_import(db, table, strategy)

    if expected_total_rows is not None and expected_total_rows != plan.total_rows:
        raise ImportConflictError(
            f"提交的行数（{expected_total_rows}）与本次解析结果（{plan.total_rows}）不一致，"
            "请重新预览后再执行",
            code="A01",
        )
    if plan.cycle:
        raise ImportCycleError(
            "层级关系存在循环，已整批拒绝，未写入任何数据：" + " → ".join(plan.cycle),
            cycle=plan.cycle,
        )
    row_errors = [
        issue
        for item in plan.rows
        for issue in item.parsed.all_issues
        if issue.level == "error"
    ]
    if row_errors:
        raise ImportConflictError(
            f"存在 {plan.error_rows} 行错误，已整批拒绝，未写入任何数据",
            code="A01",
            issues=row_errors,
        )

    try:
        batch = KnowledgePointImportBatch(
            user_id=user_id,
            source_name=source_name,
            source_format=table.source_format,
            conflict_strategy=strategy,
            status="applied",
            total_rows=plan.total_rows,
        )
        db.add(batch)
        db.flush()  # 取批次 id，供 import_batch_id 使用

        service = KnowledgePointService(db)
        specs = _creation_specs(plan)
        created_ids: dict[tuple, int] = {}
        visiting: set[tuple] = set()

        def ensure(node: tuple | None) -> int | None:
            """按需创建节点并返回其 id —— 递归保证「父级先于子级」存在。"""
            if node is None:
                return None
            if node[0] == "kp":
                return int(node[1])
            if node in created_ids:
                return created_ids[node]
            spec = specs.get(node)
            if spec is None:
                return None
            if node in visiting:
                raise ImportCycleError(
                    f"层级关系存在循环，已整批拒绝：{node[2]}", cycle=[node[2]]
                )
            visiting.add(node)
            parent_id = ensure(spec.parent_key)
            created = service.create(
                name=spec.name,
                subject=spec.subject,
                parent_id=parent_id,
                source=KP_SOURCE_ADMIN,
                import_batch_id=batch.id,
                **spec.extras,
            )
            visiting.discard(node)
            created_ids[node] = created.id
            return created.id

        auto_parent_keys = {auto.key for auto in plan.auto_parents}
        created_count = 0
        updated_count = 0
        skipped_count = 0
        for planned in plan.rows:
            row = planned.parsed
            if planned.action == "create":
                ensure(_new_node(row.normalized_subject, row.normalized_name))
                created_count += 1
            elif planned.action == "update_empty" and planned.existing is not None:
                payload = _fill_empty_payload(planned.existing, row)
                if payload is None:
                    skipped_count += 1
                else:
                    service.update(planned.existing.id, payload)
                    updated_count += 1
            else:
                skipped_count += 1

        auto_parent_count = len(auto_parent_keys & set(created_ids))
        batch.created_count = created_count
        batch.updated_count = updated_count
        batch.skipped_count = skipped_count
        batch.failed_count = 0
        batch.auto_parent_count = auto_parent_count
        batch.applied_at = datetime.now(timezone.utc)
        db.commit()
    except Exception as exc:
        # 单事务语义：任何写入期异常一律整体回滚，零写入。
        db.rollback()
        _record_failed_batch(
            db,
            user_id=user_id,
            source_name=source_name,
            source_format=table.source_format,
            strategy=strategy,
            total_rows=plan.total_rows,
            summary=f"导入失败，已整体回滚：{exc.__class__.__name__}: {exc}",
        )
        raise ImportWriteError(
            f"导入失败，已整体回滚，未写入任何数据：{exc}", code="A02"
        ) from exc

    return ApplyReport(
        batch_id=batch.id,
        total_rows=plan.total_rows,
        created_count=created_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        auto_parent_count=auto_parent_count,
        duration_ms=int((time.monotonic() - started) * 1000),
        rows=[preview_row(item) for item in plan.rows],
    )


def _fill_empty_payload(
    existing: KnowledgePoint, row: ParsedRow
) -> KnowledgePointUpdate | None:
    """`update_empty`：只补原值为空的字段，绝不覆盖已有非空值。

    以 `KnowledgePointUpdate` 交给 `KnowledgePointService.update()` —— 该方法的
    `model_fields_set` 天然等于「本次要动的字段」，因此复用了既有的写入与校验
    逻辑，而不是另写一套 UPDATE。
    """
    changes: dict[str, object] = {}
    if not existing.code and row.code:
        changes["code"] = row.code
    if not existing.aliases and row.aliases:
        changes["aliases"] = row.aliases
    if not existing.difficulty and row.difficulty:
        changes["difficulty"] = row.difficulty
    if not existing.estimated_minutes and row.estimated_minutes:
        changes["estimated_minutes"] = row.estimated_minutes
    if not changes:
        return None
    return KnowledgePointUpdate(**changes)


def _record_failed_batch(
    db: Session,
    *,
    user_id: int,
    source_name: str | None,
    source_format: str,
    strategy: str,
    total_rows: int,
    summary: str,
) -> None:
    """在被回滚的事务之外，用新事务单独记录失败批次。

    否则「批次表里看不到这次失败」—— 这正是 `status='failed'` 这个枚举值存在
    的原因（M1 已把它写进 CHECK）。`error_summary` 只存供人读的摘要，不存全量行数据。
    """
    batch = KnowledgePointImportBatch(
        user_id=user_id,
        source_name=source_name,
        source_format=source_format,
        conflict_strategy=strategy,
        status="failed",
        total_rows=total_rows,
        failed_count=total_rows,
        error_summary=summary[:2000],
    )
    db.add(batch)
    db.commit()


# --------------------------------------------------------------------------
# 回滚（L3 批次回滚）
# --------------------------------------------------------------------------


def collect_blockers(db: Session, items: list[KnowledgePoint]) -> list[BlockingRef]:
    """回滚预检：4 类引用任一命中即整批拒绝（语义与 `delete()` 完全一致）。"""
    if not items:
        return []
    ids = {item.id for item in items}
    names = {item.id: item.name for item in items}
    blockers: list[BlockingRef] = []
    seen: set[tuple[int, str]] = set()

    def _add(kp_id: int, reason: str, detail: str) -> None:
        if (kp_id, reason) in seen:
            return
        seen.add((kp_id, reason))
        blockers.append(
            BlockingRef(
                knowledge_point_id=kp_id,
                knowledge_point_name=names.get(kp_id, str(kp_id)),
                reason=reason,  # type: ignore[arg-type]
                detail=detail,
            )
        )

    # 1) 存在子知识点，且该子节点不属于本批（同批父子会按叶子到根一起删）。
    for child_id, parent_id in db.execute(
        select(KnowledgePoint.id, KnowledgePoint.parent_id).where(
            KnowledgePoint.parent_id.in_(ids)
        )
    ).all():
        if child_id not in ids:
            _add(int(parent_id), "has_children", f"子知识点「id={child_id}」不属于本批次")

    # 2) 已被题目关联。
    for kp_id in db.scalars(
        select(QuestionKnowledgePoint.knowledge_point_id)
        .where(QuestionKnowledgePoint.knowledge_point_id.in_(ids))
        .distinct()
    ).all():
        _add(int(kp_id), "linked_question", "已与题目建立关联，请先解除关联")

    # 3) 已有用户掌握度记录。
    for kp_id in db.scalars(
        select(UserKnowledgePointMastery.knowledge_point_id)
        .where(UserKnowledgePointMastery.knowledge_point_id.in_(ids))
        .distinct()
    ).all():
        _add(int(kp_id), "has_mastery", "已有用户掌握度记录，无法删除")

    # 4) 已建立前置依赖边（任一方向）。
    for kp_id, prerequisite_id in db.execute(
        select(
            KnowledgePointPrerequisite.knowledge_point_id,
            KnowledgePointPrerequisite.prerequisite_id,
        ).where(
            KnowledgePointPrerequisite.knowledge_point_id.in_(ids)
            | KnowledgePointPrerequisite.prerequisite_id.in_(ids)
        )
    ).all():
        target = int(kp_id) if kp_id in ids else int(prerequisite_id)
        _add(target, "prerequisite_edge", "已建立前置依赖关系，请先解除前置关系")

    return blockers


@dataclass
class RollbackResult:
    batch_id: int
    deleted_count: int
    kept_updated_count: int
    rolled_back_at: datetime


def rollback_batch(db: Session, batch_id: int) -> RollbackResult:
    batch = db.get(KnowledgePointImportBatch, batch_id)
    if batch is None:
        raise ImportBatchNotFound("批次不存在")
    if batch.status == "rolled_back":
        raise ImportConflictError("该批次已经回滚过，无法重复回滚", code="A03")
    if batch.status != "applied":
        raise ImportConflictError("失败批次没有可回滚的数据", code="A03")

    items = list(
        db.scalars(
            select(KnowledgePoint).where(KnowledgePoint.import_batch_id == batch_id)
        ).all()
    )
    blockers = collect_blockers(db, items)
    if blockers:
        raise ImportConflictError(
            f"有 {len(blockers)} 处引用阻止回滚，未删除任何数据",
            code="A03",
            issues=[
                RowIssue(
                    row=None,
                    level="error",
                    code="A03",
                    message=f"知识点「{item.knowledge_point_name}」：{item.detail}",
                )
                for item in blockers
            ],
        )

    deleted = _delete_leaf_to_root(db, {item.id for item in items})
    now = datetime.now(timezone.utc)
    batch.status = "rolled_back"
    batch.rolled_back_at = now
    # 计数列保持原值：与 status 一起构成「本批曾导入了多少」的完整审计事实。
    kept_updated_count = batch.updated_count
    db.commit()
    return RollbackResult(
        batch_id=batch_id,
        deleted_count=deleted,
        kept_updated_count=kept_updated_count,
        rolled_back_at=now,
    )


def _delete_leaf_to_root(db: Session, ids: set[int]) -> int:
    """`parent_id` 是 ON DELETE RESTRICT，因此必须从叶子往根逐层删。

    用 Core DELETE（而不是 `db.delete(obj)`）是刻意的：ORM 的 unit-of-work 会在
    删除父节点时反向把子节点的外键置空，与「按层删除」的语义冲突；Core DELETE
    完全绕开关系级联，行为可预测。
    """
    if not ids:
        return 0
    parent_of = dict(
        db.execute(
            select(KnowledgePoint.id, KnowledgePoint.parent_id).where(
                KnowledgePoint.id.in_(ids)
            )
        ).all()
    )
    remaining_children = {kp_id: 0 for kp_id in parent_of}
    for parent_id in parent_of.values():
        if parent_id in remaining_children:
            remaining_children[parent_id] += 1
    frontier = [kp_id for kp_id, count in remaining_children.items() if count == 0]
    deleted = 0
    while frontier:
        db.execute(sa_delete(KnowledgePoint).where(KnowledgePoint.id.in_(frontier)))
        deleted += len(frontier)
        nxt: list[int] = []
        for kp_id in frontier:
            parent_id = parent_of.get(kp_id)
            if parent_id in remaining_children:
                remaining_children[parent_id] -= 1
                if remaining_children[parent_id] == 0:
                    nxt.append(parent_id)
        frontier = nxt
    if deleted != len(parent_of):
        raise ImportConflictError(
            "本批次内部存在父子环，无法安全删除，已停止回滚", code="A03"
        )
    return deleted


# --------------------------------------------------------------------------
# 模板
# --------------------------------------------------------------------------

TEMPLATE_HEADER = "学科,层级路径,知识点名,别名,难度,预计学时,编码"
# 样例刻意让**层级路径的每一段都只作目录**（声明为知识点的是「知识点名」列），
# 因此不会触发 E12「同名同时被作为目录与知识点」的类型冲突。
TEMPLATE_SAMPLE_ROWS = (
    "数据结构,线性表,顺序表,linear list|sequential list,easy,30,DS.LINEAR.SEQUENTIAL",
    "数据结构,线性表>链表,单链表,singly linked list|linked list,medium,40,"
    "DS.LINEAR.LINKED_LIST.SINGLE",
    "数据结构,树与二叉树>二叉树遍历,先序遍历,preorder traversal,medium,40,"
    "DS.TREE.BINARY.TRAVERSAL.PREORDER",
)


def template_csv() -> str:
    """CSV 模板正文（不含 BOM —— 由响应层负责加，避免重复）。"""
    return "\r\n".join([TEMPLATE_HEADER, *TEMPLATE_SAMPLE_ROWS]) + "\r\n"
